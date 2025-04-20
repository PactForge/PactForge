import os
import time
import chromadb
from flask import Flask, request, jsonify, render_template
from docx import Document
import google.generativeai as genai
from google.api_core import retry
from google.generativeai.types import GenerationConfig

app = Flask(__name__)

# Configure Google Gemini API
GOOGLE_API_KEY = os.getenv("GOOGLE_API_KEY")
if not GOOGLE_API_KEY:
    raise ValueError("GOOGLE_API_KEY environment variable not set")
genai.configure(api_key=GOOGLE_API_KEY)

# ChromaDB setup
clientdb = chromadb.PersistentClient(path="./chroma_db")
rent = clientdb.get_or_create_collection(name="rent_agreements")
nda = clientdb.get_or_create_collection(name="nda_agreements")
employment = clientdb.get_or_create_collection(name="employ_agreements")
franchise = clientdb.get_or_create_collection(name="franchise_agreements")
contractor = clientdb.get_or_create_collection(name="contractor_agreements")
all_dbs = [rent, nda, employment, franchise, contractor]

# Model configuration
model_config = GenerationConfig(temperature=0.75, top_p=0.9)

# Retry mechanism for API calls
def is_retriable(e):
    return isinstance(e, Exception) and hasattr(e, 'code') and e.code in {429, 503}

@retry.Retry(predicate=is_retriable)
def generate_embeddings(cl, etype):
    embedding_task = "retrieval_document" if etype else "retrieval_query"
    response = genai.embed_content(
        model="models/text-embedding-004",
        content=cl,
        task_type=embedding_task
    )
    return [response['embedding']]

def read_docx(endname):
    path = f"./Clauses/{endname}.docx"
    if not os.path.exists(path):
        print(f"Warning: {path} not found, returning empty clauses")
        return []
    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return paragraphs

def extract_doxc(path):
    if not os.path.exists(path):
        print(f"Warning: {path} not found, returning None")
        return None
    doc = Document(path)
    return doc

def extract_samples(endname):
    dataset_path = f"./sampleagreements/{endname}/{endname}"
    if not os.path.exists(dataset_path):
        print(f"Warning: {dataset_path} not found, returning empty list")
        return []
    docx_files = []
    for item in os.listdir(dataset_path):
        item_path = os.path.join(dataset_path, item)
        if os.path.isfile(item_path) and item.lower().endswith('.docx'):
            docx_files.append(item_path)
    return docx_files

# Initialize ChromaDB with clauses
def initialize_chromadb():
    agreement_types = ["rent", "nda", "employment", "franchise", "contractor"]
    all_clauses = []
    for t in agreement_types:
        clauses = read_docx(t)
        all_clauses.append(clauses)

    for j, dataset in enumerate(all_clauses):
        if not dataset:
            print(f"No clauses for {agreement_types[j]}, collection will be empty")
            continue
        embeds, ids, documents = [], [], []
        for i, clause in enumerate(dataset):
            try:
                vector = generate_embeddings(clause, True)
                time.sleep(0.4)  # Rate limit
                embeds.append(vector[0])
                ids.append(f"clause-{j}-{i}")
                documents.append(clause)
            except Exception as e:
                print(f"Error embedding clause {i} for {agreement_types[j]}: {e}")
        if embeds:
            all_dbs[j].add(embeddings=embeds, ids=ids, documents=documents)
            print(f"Initialized {agreement_types[j]} with {len(embeds)} clauses")

# Run initialization if collections are empty
if not any(db.count() > 0 for db in all_dbs):
    initialize_chromadb()

def strip_type(agr: str):
    agreement_types = ["rent", "nda", "contractor", "employment", "franchise"]
    model = genai.GenerativeModel('gemini-1.5-flash')
    prompt = f"""Return the type of agreement in one lowercase word from {agreement_types}. Input: {agr}"""
    response = model.generate_content(prompt, generation_config=model_config)
    return response.text.strip().lower()

def pos_neg(response: str):
    model = genai.GenerativeModel('gemini-1.5-flash')
    prompt = f"""Classify sentiment: Reply '1' for positive, '0' for negative. Sentence = {response}"""
    response_heat = model.generate_content(prompt, generation_config=model_config)
    return bool(int(response_heat.text))

def perform_analysis(atype, impt):
    model = genai.GenerativeModel('gemini-1.5-flash')
    prompt = f"""As a legal assistant, evaluate if {impt} is sufficient for a {atype} agreement. Respond with:
    - 'Yes. All essential information seems to be present.' if comprehensive.
    - 'No, The following essential information seems to be missing or unclear: [list]' if details are missing.
    - 'No, The provided information is too vague or insufficient.' if vague."""
    response = model.generate_content(prompt, generation_config=model_config)
    return response.text, pos_neg(response.text)

def obtain_information_holes(final_type, important_info, extra_info):
    model = genai.GenerativeModel('gemini-1.5-flash')
    total_info = important_info + extra_info
    prompt = f"""Identify missing or unclear information for a {final_type} agreement based on: {total_info}. Return specific details needed."""
    response = model.generate_content(prompt, generation_config=model_config)
    return response.text

def get_data(holes: str, final_type):
    model = genai.GenerativeModel('gemini-1.5-flash')
    prompt = f"""Retrieve information for {holes} to generate a {final_type} agreement."""
    response = model.generate_content(prompt, generation_config=model_config)
    return response.text

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/send_message', methods=['POST'])
def send_message():
    global final_type, important_info, extra_info, req
    data = request.json
    user_input = data.get('message', '').strip()

    if not hasattr(app, 'state'):
        app.state = {
            'req': False,
            'final_type': None,
            'important_info': '',
            'extra_info': '',
            'step': 'agreement_type'
        }

    state = app.state

    if state['step'] == 'agreement_type':
        final_type = strip_type(user_input)
        agreement_types = ["rent", "nda", "contractor", "employment", "franchise"]
        if final_type not in agreement_types:
            return jsonify({'response': "Invalid agreement type. Please choose: rent, nda, contractor, employment, franchise"})
        state['final_type'] = final_type
        state['step'] = 'important_info'
        return jsonify({'response': f"Please provide important information for your {final_type} agreement (e.g., parties, duration, financial details)."})

    elif state['step'] == 'important_info':
        state['important_info'] = user_input
        state['step'] = 'extra_info'
        return jsonify({'response': f"Provide any extra information to tailor your {state['final_type']} agreement (e.g., specific clauses, conditions)."})

    elif state['step'] == 'extra_info':
        state['extra_info'] = user_input
        analysis, is_positive = perform_analysis(state['final_type'], state['important_info'])
        if not is_positive:
            state['step'] = 'important_info'
            return jsonify({'response': analysis + " Please provide more information."})

        state['req'] = True
        # Retrieve relevant clauses
        dbname = state['final_type'] + "_agreements"
        querydb = clientdb.get_collection(name=dbname)
        user_query = state['important_info'] + state['extra_info']
        query_embed = generate_embeddings(user_query, False)
        results = querydb.query(query_embeddings=query_embed, n_results=querydb.count())
        relevant_documents = results['documents'] if results['documents'] else []

        # Get sample agreements
        sample_agreements = extract_samples(state['final_type'])

        # Get additional info
        info_holes = obtain_information_holes(state['final_type'], state['important_info'], state['extra_info'])
        obtained_info = get_data(info_holes, state['final_type'])

        # Generate final agreement
        model = genai.GenerativeModel('gemini-1.5-flash')
        prompt = f"""Generate a {state['final_type']} agreement using:
        User info: {state['important_info']}
        Extra info: {state['extra_info']}
        Relevant clauses: {relevant_documents}
        Sample agreements: {sample_agreements}
        Additional info: {obtained_info}
        Ensure the agreement is concise, legally robust, and clear."""

        response = model.generate_content(prompt, generation_config=model_config)

        # Reset state
        app.state = {'req': False, 'final_type': None, 'important_info': '', 'extra_info': '', 'step': 'agreement_type'}

        return jsonify({'response': response.text})

# The following lines should be removed or commented out for Render deployment
# port = int(os.environ.get('PORT', 5000))
